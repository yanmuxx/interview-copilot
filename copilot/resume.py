"""简历解析：从 PDF/DOCX/TXT/MD 提取文字，并用 LLM 提炼成"候选人画像"。"""
from __future__ import annotations

import io

MAX_RESUME_CHARS = 6000

DISTILL_PROMPT = """\
你是简历分析助手。请把下面的简历提炼成给"实时面试答题助手"使用的候选人画像，
用 Markdown 输出，控制在 300 字以内，包含这些小节：

## 基本信息
（工作年限、方向、当前角色）

## 技术亮点
（3~5 条，每条一句话）

## 项目亮点
（2~3 条，优先带数字/结果的项目）

## 求职意向
（目标岗位/方向；简历没有就从内容推断）

只输出画像本身，不要任何解释。画像将作为答题提示的背景资料，要突出能用来"圆场"的优势点。

【岗位 JD】
{jd}

【简历全文】
{resume}
"""


def extract_text(filename: str, data: bytes) -> str:
    suffix = (filename or "").lower().rsplit(".", 1)[-1]
    if suffix in ("txt", "md", "markdown"):
        return data.decode("utf-8", errors="ignore")
    if suffix == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError("缺少 pypdf，请 pip install pypdf") from e
        pages = PdfReader(io.BytesIO(data)).pages
        text = "\n".join((page.extract_text() or "") for page in pages).strip()
        if not text:
            raise RuntimeError("这个 PDF 提取不出文字（可能是扫描版/图片版），请转成 docx 或 txt 再上传")
        return text
    if suffix == "docx":
        try:
            from docx import Document
        except ImportError as e:
            raise RuntimeError("缺少 python-docx，请 pip install python-docx") from e
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(t for t in parts if t.strip())
    raise RuntimeError(f"不支持的文件类型 .{suffix}，请上传 PDF / DOCX / TXT / MD")


def distill_persona(resume_text: str, jd: str, api_key: str, base_url: str, model: str) -> str:
    from openai import OpenAI

    if not api_key:
        raise RuntimeError("请先在「LLM 设置」里填入 API Key，才能用 AI 解析简历")
    client = OpenAI(api_key=api_key, base_url=base_url, timeout=60, max_retries=1)
    resp = client.chat.completions.create(
        model=model,
        messages=[{
            "role": "user",
            "content": DISTILL_PROMPT.format(
                jd=(jd or "（未提供）")[:1500],
                resume=resume_text[:MAX_RESUME_CHARS],
            ),
        }],
        temperature=0.2,
        max_tokens=2048,  # 思考型模型的思考过程也计入配额，600 会全部被思考吃掉导致正文为空
    )
    content = (resp.choices[0].message.content or "").strip()
    if not content:
        raise RuntimeError(
            "模型没有返回正文（输出配额被思考过程耗尽）。请在设置里换非思考模型（如 deepseek-chat），或稍后重试"
        )
    return content
