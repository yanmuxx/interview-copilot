"""简历文本抽取测试。"""
import pytest

from copilot.resume import extract_text


def test_txt_extraction(tmp_path):
    f = tmp_path / "r.txt"
    f.write_text("张三，5年后端经验", encoding="utf-8")
    assert "后端" in extract_text("r.txt", f.read_bytes())


def test_md_extraction(tmp_path):
    f = tmp_path / "r.md"
    f.write_text("# 简历\n- 熟悉 Redis", encoding="utf-8")
    assert "Redis" in extract_text("r.md", f.read_bytes())


def test_docx_extraction(tmp_path):
    from docx import Document

    f = tmp_path / "r.docx"
    doc = Document()
    doc.add_paragraph("李四，红队方向，3年经验")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "证书"
    t.rows[0].cells[1].text = "CISP-PTE"
    f.write_bytes(b"")  # 占位，内容由 Document.save 写入
    doc.save(str(f))
    text = extract_text("r.docx", f.read_bytes())
    assert "红队" in text and "CISP-PTE" in text  # 段落和表格都要抽出来


def test_unsupported_suffix_raises():
    with pytest.raises(RuntimeError):
        extract_text("r.exe", b"MZ")


def test_empty_pdf_raises_clear_error():
    # 构造一个无文字层的"伪 PDF"（空页）触发明确报错
    with pytest.raises(RuntimeError, match="扫描版|提取不出|解析失败"):
        extract_text("r.pdf", b"%PDF-1.4\n%%EOF\n")
